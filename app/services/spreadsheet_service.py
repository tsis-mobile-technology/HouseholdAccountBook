import io
import json
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.connection import get_connection, get_db_cursor
from app.services.analytics_service import get_annual_dashboard, get_monthly_summary
from app.services.fixed_plan_service import get_month_fixed_plans
from app.services.transaction_service import get_variable_expenses

# Style definitions
FONT_TITLE = Font(name="맑은 고딕", size=15, bold=True, color="1F2937")
FONT_SECTION = Font(name="맑은 고딕", size=11, bold=True, color="1F2937")
FONT_HEADER = Font(name="맑은 고딕", size=10, bold=True, color="374151")
FONT_BODY = Font(name="맑은 고딕", size=10, color="1F2937")
FONT_BOLD = Font(name="맑은 고딕", size=10, bold=True, color="1F2937")

FILL_HEADER = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
FILL_ACCENT = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # soft blue
FILL_GREEN = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # soft green
FILL_PURPLE = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid") # soft purple

THIN_BORDER = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

def export_to_excel_bytes(year: int) -> bytes:
    """Generates an Excel workbook with Annual Dashboard + 12 Monthly sheets matching Google Sheet."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Sheet: 연간 대시보드
    ws_ann = wb.create_sheet(title="연간 대시보드")
    annual_data = get_annual_dashboard(year)

    ws_ann.merge_cells("B2:I2")
    ws_ann["B2"] = f"{year}년 연간 가계부 대시보드"
    ws_ann["B2"].font = FONT_TITLE
    ws_ann["B2"].alignment = Alignment(vertical="center")

    headers_ann = ["월", "총 수입(원)", "고정 지출(원)", "변동 지출(원)", "총 지출(원)", "저축/투자(원)", "당월 잔액(원)", "저축률"]
    for col_idx, h in enumerate(headers_ann, start=2):
        cell = ws_ann.cell(row=4, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    current_row = 5
    for m in annual_data.months:
        ws_ann.cell(row=current_row, column=2, value=m.month_name).alignment = Alignment(horizontal="center")
        ws_ann.cell(row=current_row, column=3, value=m.total_income).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=4, value=m.fixed_expense).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=5, value=m.variable_expense).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=6, value=m.total_expense).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=7, value=m.savings_investment).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=8, value=m.balance).number_format = "#,##0"
        ws_ann.cell(row=current_row, column=9, value=m.savings_rate / 100.0).number_format = "0.0%"

        for c in range(2, 10):
            ws_ann.cell(row=current_row, column=c).font = FONT_BODY
            ws_ann.cell(row=current_row, column=c).border = THIN_BORDER
        current_row += 1

    # 연간 합계
    tot = annual_data.annual_total
    ws_ann.cell(row=current_row, column=2, value=tot.month_name).alignment = Alignment(horizontal="center")
    ws_ann.cell(row=current_row, column=3, value=tot.total_income).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=4, value=tot.fixed_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=5, value=tot.variable_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=6, value=tot.total_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=7, value=tot.savings_investment).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=8, value=tot.balance).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=9, value=tot.savings_rate / 100.0).number_format = "0.0%"
    for c in range(2, 10):
        c_cell = ws_ann.cell(row=current_row, column=c)
        c_cell.font = FONT_BOLD
        c_cell.fill = FILL_ACCENT
        c_cell.border = THIN_BORDER
    current_row += 1

    # 월평균
    avg = annual_data.monthly_average
    ws_ann.cell(row=current_row, column=2, value=avg.month_name).alignment = Alignment(horizontal="center")
    ws_ann.cell(row=current_row, column=3, value=avg.total_income).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=4, value=avg.fixed_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=5, value=avg.variable_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=6, value=avg.total_expense).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=7, value=avg.savings_investment).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=8, value=avg.balance).number_format = "#,##0"
    ws_ann.cell(row=current_row, column=9, value=avg.savings_rate / 100.0).number_format = "0.0%"
    for c in range(2, 10):
        c_cell = ws_ann.cell(row=current_row, column=c)
        c_cell.font = FONT_BOLD
        c_cell.fill = FILL_HEADER
        c_cell.border = THIN_BORDER

    # Set column widths
    for col in ws_ann.columns:
        col_letter = get_column_letter(col[0].column)
        ws_ann.column_dimensions[col_letter].width = 15
    ws_ann.column_dimensions['A'].width = 3

    # 2. Sheets: 1월 ~ 12월
    for m in range(1, 13):
        ws_m = wb.create_sheet(title=f"{m}월")
        ws_m.cell(row=2, column=2, value=f"{year}년 {m}월 가계부").font = FONT_TITLE

        # Left Section: Fixed Plans
        ws_m.cell(row=4, column=2, value="[고정 수입]").font = FONT_SECTION
        for i, h in enumerate(["항목", "예정일", "내용", "금액(원)"], start=2):
            c = ws_m.cell(row=5, column=i, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_HEADER
            c.border = THIN_BORDER

        plans = get_month_fixed_plans(year, m)
        r_idx = 6
        for item in plans["incomes"]:
            ws_m.cell(row=r_idx, column=2, value=item.item_name).font = FONT_BODY
            ws_m.cell(row=r_idx, column=3, value=item.day).font = FONT_BODY
            ws_m.cell(row=r_idx, column=4, value=item.description).font = FONT_BODY
            c_amt = ws_m.cell(row=r_idx, column=5, value=item.amount)
            c_amt.font = FONT_BODY
            c_amt.number_format = "#,##0"
            for c in range(2, 6):
                ws_m.cell(row=r_idx, column=c).border = THIN_BORDER
            r_idx += 1

        # Fixed Income Total
        ws_m.cell(row=r_idx, column=2, value="고정수입 합계").font = FONT_BOLD
        c_tot = ws_m.cell(row=r_idx, column=5, value=sum(i.amount for i in plans["incomes"]))
        c_tot.font = FONT_BOLD
        c_tot.number_format = "#,##0"
        for c in range(2, 6):
            ws_m.cell(row=r_idx, column=c).fill = FILL_GREEN
            ws_m.cell(row=r_idx, column=c).border = THIN_BORDER

        # Fixed Expenses
        r_idx += 2
        ws_m.cell(row=r_idx, column=2, value="[고정 지출]").font = FONT_SECTION
        r_idx += 1
        for i, h in enumerate(["항목", "출금일", "내용/기관", "금액(원)"], start=2):
            c = ws_m.cell(row=r_idx, column=i, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_HEADER
            c.border = THIN_BORDER
        r_idx += 1
        for item in plans["expenses"]:
            ws_m.cell(row=r_idx, column=2, value=item.item_name).font = FONT_BODY
            ws_m.cell(row=r_idx, column=3, value=item.day).font = FONT_BODY
            ws_m.cell(row=r_idx, column=4, value=item.description).font = FONT_BODY
            c_amt = ws_m.cell(row=r_idx, column=5, value=item.amount)
            c_amt.font = FONT_BODY
            c_amt.number_format = "#,##0"
            for c in range(2, 6):
                ws_m.cell(row=r_idx, column=c).border = THIN_BORDER
            r_idx += 1

        # Fixed Expense Total
        ws_m.cell(row=r_idx, column=2, value="고정지출 합계").font = FONT_BOLD
        c_tot = ws_m.cell(row=r_idx, column=5, value=sum(e.amount for e in plans["expenses"]))
        c_tot.font = FONT_BOLD
        c_tot.number_format = "#,##0"
        for c in range(2, 6):
            ws_m.cell(row=r_idx, column=c).fill = FILL_PURPLE
            ws_m.cell(row=r_idx, column=c).border = THIN_BORDER

        # Savings & Investment
        r_idx += 2
        ws_m.cell(row=r_idx, column=2, value="[저축 및 투자]").font = FONT_SECTION
        r_idx += 1
        for i, h in enumerate(["항목", "납입일", "금융기관/내용", "금액(원)"], start=2):
            c = ws_m.cell(row=r_idx, column=i, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_HEADER
            c.border = THIN_BORDER
        r_idx += 1
        for item in plans["savings"]:
            ws_m.cell(row=r_idx, column=2, value=item.item_name).font = FONT_BODY
            ws_m.cell(row=r_idx, column=3, value=item.day).font = FONT_BODY
            ws_m.cell(row=r_idx, column=4, value=item.description).font = FONT_BODY
            c_amt = ws_m.cell(row=r_idx, column=5, value=item.amount)
            c_amt.font = FONT_BODY
            c_amt.number_format = "#,##0"
            for c in range(2, 6):
                ws_m.cell(row=r_idx, column=c).border = THIN_BORDER
            r_idx += 1

        # Savings Total
        ws_m.cell(row=r_idx, column=2, value="저축/투자 합계").font = FONT_BOLD
        c_tot = ws_m.cell(row=r_idx, column=5, value=sum(s.amount for s in plans["savings"]))
        c_tot.font = FONT_BOLD
        c_tot.number_format = "#,##0"
        for c in range(2, 6):
            ws_m.cell(row=r_idx, column=c).fill = FILL_ACCENT
            ws_m.cell(row=r_idx, column=c).border = THIN_BORDER

        # Summary box
        r_idx += 2
        summary = get_monthly_summary(year, m)
        ws_m.cell(row=r_idx, column=2, value=f"[{m}월 재무 요약]").font = FONT_SECTION
        sum_rows = [
            ("총 수입", summary.total_income),
            ("고정 지출", summary.fixed_expense),
            ("변동 지출", summary.variable_expense),
            ("총 지출", summary.total_expense),
            ("저축 및 투자", summary.savings_investment),
            ("당월 잔액(잉여금)", summary.balance),
        ]
        for s_title, s_val in sum_rows:
            r_idx += 1
            ws_m.cell(row=r_idx, column=2, value=s_title).font = FONT_BODY
            c_val = ws_m.cell(row=r_idx, column=3, value=s_val)
            c_val.font = FONT_BOLD
            c_val.number_format = "#,##0"
            ws_m.cell(row=r_idx, column=2).border = THIN_BORDER
            c_val.border = THIN_BORDER

        # Right Section: [이달의 변동 지출 내역] (Columns G..L)
        ws_m.cell(row=4, column=7, value="[이달의 변동 지출 내역]").font = FONT_SECTION
        v_headers = ["일자", "분류", "지출내용", "결제수단", "금액(원)", "비고"]
        for i, h in enumerate(v_headers, start=7):
            c = ws_m.cell(row=5, column=i, value=h)
            c.font = FONT_HEADER
            c.fill = FILL_HEADER
            c.border = THIN_BORDER

        v_expenses = get_variable_expenses(year=year, month=m, limit=1000)
        # Reverse to chronological order for excel sheet
        v_expenses_sorted = sorted(v_expenses, key=lambda x: (x.transaction_date, x.id))

        vr_idx = 6
        for ve in v_expenses_sorted:
            ws_m.cell(row=vr_idx, column=7, value=ve.transaction_date).alignment = Alignment(horizontal="center")
            ws_m.cell(row=vr_idx, column=8, value=ve.category_name).alignment = Alignment(horizontal="center")
            ws_m.cell(row=vr_idx, column=9, value=ve.title)
            ws_m.cell(row=vr_idx, column=10, value=ve.payment_method_name).alignment = Alignment(horizontal="center")
            c_amt = ws_m.cell(row=vr_idx, column=11, value=ve.amount)
            c_amt.number_format = "#,##0"
            ws_m.cell(row=vr_idx, column=12, value=ve.memo or "")
            for c in range(7, 13):
                ws_m.cell(row=vr_idx, column=c).font = FONT_BODY
                ws_m.cell(row=vr_idx, column=c).border = THIN_BORDER
            vr_idx += 1

        # Column widths
        ws_m.column_dimensions['A'].width = 3
        ws_m.column_dimensions['B'].width = 16
        ws_m.column_dimensions['C'].width = 14
        ws_m.column_dimensions['D'].width = 20
        ws_m.column_dimensions['E'].width = 15
        ws_m.column_dimensions['F'].width = 4
        ws_m.column_dimensions['G'].width = 13
        ws_m.column_dimensions['H'].width = 14
        ws_m.column_dimensions['I'].width = 24
        ws_m.column_dimensions['J'].width = 14
        ws_m.column_dimensions['K'].width = 15
        ws_m.column_dimensions['L'].width = 20

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

def import_from_excel_file(file_bytes: bytes) -> Dict[str, Any]:
    """Parses an uploaded XLSX file matching the Google Sheet format and saves into DB."""
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    imported_counts = {"variable_expenses": 0, "fixed_incomes": 0, "fixed_expenses": 0, "savings": 0}

    # Helper maps
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id, name FROM categories")
        cat_map = {r["name"]: r["id"] for r in cur.fetchall()}

        cur.execute("SELECT id, name FROM payment_methods")
        pay_map = {r["name"]: r["id"] for r in cur.fetchall()}
    finally:
        conn.close()

    def get_or_create_cat(name: str) -> int:
        if name in cat_map:
            return cat_map[name]
        with get_db_cursor() as cur:
            cur.execute("INSERT OR IGNORE INTO categories (name, color_hex, icon) VALUES (?, '#BEE3F8', 'tag')", (name,))
            cur.execute("SELECT id FROM categories WHERE name = ?", (name,))
            c_id = cur.fetchone()["id"]
            cat_map[name] = c_id
            return c_id

    def get_or_create_pay(name: str) -> int:
        if name in pay_map:
            return pay_map[name]
        with get_db_cursor() as cur:
            cur.execute("INSERT OR IGNORE INTO payment_methods (name) VALUES (?)", (name,))
            cur.execute("SELECT id FROM payment_methods WHERE name = ?", (name,))
            p_id = cur.fetchone()["id"]
            pay_map[name] = p_id
            return p_id

    # Iterate through sheets
    for sheetname in wb.sheetnames:
        # Check if sheet is like "1월", "2월"...
        if not sheetname.endswith("월") or sheetname == "연간 대시보드":
            continue

        try:
            m_num = int(sheetname.replace("월", "").strip())
        except ValueError:
            continue

        ws = wb[sheetname]
        # Detect year from B2 (e.g. "2027년 1월 가계부")
        b2_val = str(ws["B2"].value or "")
        year = datetime.now().year
        if "년" in b2_val:
            try:
                year = int(b2_val.split("년")[0].strip())
            except Exception:
                pass

        # 1. Fixed Incomes (Rows 6..9 roughly, until "고정수입 합계")
        incomes = []
        for r in range(6, 12):
            item = ws.cell(row=r, column=2).value
            if not item or str(item).strip() == "고정수입 합계":
                break
            day = str(ws.cell(row=r, column=3).value or "")
            desc = str(ws.cell(row=r, column=4).value or "")
            amt = int(float(ws.cell(row=r, column=5).value or 0))
            incomes.append((year, m_num, str(item).strip(), day, desc, amt))

        # 2. Fixed Expenses
        expenses = []
        fe_start = None
        for r in range(12, 20):
            if str(ws.cell(row=r, column=2).value or "").strip() == "[고정 지출]":
                fe_start = r + 2  # data starts after header
                break
        if fe_start:
            for r in range(fe_start, fe_start + 10):
                item = ws.cell(row=r, column=2).value
                if not item or str(item).strip() == "고정지출 합계":
                    break
                day = str(ws.cell(row=r, column=3).value or "")
                desc = str(ws.cell(row=r, column=4).value or "")
                amt = int(float(ws.cell(row=r, column=5).value or 0))
                expenses.append((year, m_num, str(item).strip(), day, desc, amt))

        # 3. Savings & Investments
        savings = []
        si_start = None
        for r in range(20, 30):
            if str(ws.cell(row=r, column=2).value or "").strip() == "[저축 및 투자]":
                si_start = r + 2
                break
        if si_start:
            for r in range(si_start, si_start + 10):
                item = ws.cell(row=r, column=2).value
                if not item or str(item).strip() == "저축/투자 합계":
                    break
                day = str(ws.cell(row=r, column=3).value or "")
                desc = str(ws.cell(row=r, column=4).value or "")
                amt = int(float(ws.cell(row=r, column=5).value or 0))
                savings.append((year, m_num, str(item).strip(), day, desc, amt))

        # Save fixed plans into DB
        with get_db_cursor() as cur:
            cur.execute("DELETE FROM fixed_incomes WHERE year = ? AND month = ?", (year, m_num))
            for inc in incomes:
                cur.execute("""
                    INSERT INTO fixed_incomes (year, month, item_name, expected_day, description, amount, is_template)
                    VALUES (?, ?, ?, ?, ?, ?, 0)
                """, inc)
                imported_counts["fixed_incomes"] += 1

            cur.execute("DELETE FROM fixed_expenses WHERE year = ? AND month = ?", (year, m_num))
            for exp in expenses:
                cur.execute("""
                    INSERT INTO fixed_expenses (year, month, item_name, withdrawal_day, description, amount, is_template)
                    VALUES (?, ?, ?, ?, ?, ?, 0)
                """, exp)
                imported_counts["fixed_expenses"] += 1

            cur.execute("DELETE FROM savings_investments WHERE year = ? AND month = ?", (year, m_num))
            for sav in savings:
                cur.execute("""
                    INSERT INTO savings_investments (year, month, item_name, payment_day, description, amount, is_template)
                    VALUES (?, ?, ?, ?, ?, ?, 0)
                """, sav)
                imported_counts["savings"] += 1

        # 4. Variable Expenses (Columns G..L, Row 6 upwards)
        var_items = []
        for r in range(6, ws.max_row + 1):
            date_val = ws.cell(row=r, column=7).value
            if not date_val:
                continue
            cat_name = str(ws.cell(row=r, column=8).value or "기타").strip()
            title = str(ws.cell(row=r, column=9).value or "").strip()
            pay_name = str(ws.cell(row=r, column=10).value or "신용카드").strip()
            amt_val = ws.cell(row=r, column=11).value
            memo_val = str(ws.cell(row=r, column=12).value or "").strip()

            if not title and not amt_val:
                continue

            try:
                amt = int(float(amt_val or 0))
            except ValueError:
                continue

            # Format date string YYYY-MM-DD
            if isinstance(date_val, datetime):
                date_str = date_val.strftime("%Y-%m-%d")
            else:
                date_str = str(date_val).strip()

            c_id = get_or_create_cat(cat_name)
            p_id = get_or_create_pay(pay_name)
            var_items.append((date_str, c_id, title, p_id, amt, memo_val))

        with get_db_cursor() as cur:
            for item in var_items:
                cur.execute("""
                    INSERT INTO variable_expenses (transaction_date, category_id, title, payment_method_id, amount, memo)
                    VALUES (?, ?, ?, ?, ?, ?)
                """, item)
                imported_counts["variable_expenses"] += 1

    return imported_counts

def export_all_json() -> Dict[str, Any]:
    """Dumps all DB content as JSON for universal portable backup."""
    conn = get_connection()
    try:
        cur = conn.cursor()
        data = {
            "version": "1.0",
            "exported_at": datetime.now().isoformat(),
            "categories": [dict(r) for r in cur.execute("SELECT * FROM categories").fetchall()],
            "payment_methods": [dict(r) for r in cur.execute("SELECT * FROM payment_methods").fetchall()],
            "fixed_incomes": [dict(r) for r in cur.execute("SELECT * FROM fixed_incomes").fetchall()],
            "fixed_expenses": [dict(r) for r in cur.execute("SELECT * FROM fixed_expenses").fetchall()],
            "savings_investments": [dict(r) for r in cur.execute("SELECT * FROM savings_investments").fetchall()],
            "variable_expenses": [dict(r) for r in cur.execute("SELECT * FROM variable_expenses").fetchall()],
        }
        return data
    finally:
        conn.close()

def restore_all_json(data: Dict[str, Any]) -> bool:
    """Restores full database from JSON export."""
    with get_db_cursor() as cur:
        # Clear tables
        cur.execute("DELETE FROM variable_expenses")
        cur.execute("DELETE FROM fixed_incomes")
        cur.execute("DELETE FROM fixed_expenses")
        cur.execute("DELETE FROM savings_investments")
        cur.execute("DELETE FROM categories")
        cur.execute("DELETE FROM payment_methods")

        for c in data.get("categories", []):
            cur.execute("""
                INSERT INTO categories (id, name, color_hex, icon, sort_order, is_active)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (c["id"], c["name"], c.get("color_hex", "#A8E6CF"), c.get("icon", "tag"), c.get("sort_order", 0), c.get("is_active", 1)))

        for p in data.get("payment_methods", []):
            cur.execute("""
                INSERT INTO payment_methods (id, name, is_active)
                VALUES (?, ?, ?)
            """, (p["id"], p["name"], p.get("is_active", 1)))

        for f in data.get("fixed_incomes", []):
            cur.execute("""
                INSERT INTO fixed_incomes (id, year, month, item_name, expected_day, description, amount, is_template)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (f.get("id"), f["year"], f["month"], f["item_name"], f.get("expected_day", ""), f.get("description", ""), f["amount"], f.get("is_template", 0)))

        for e in data.get("fixed_expenses", []):
            cur.execute("""
                INSERT INTO fixed_expenses (id, year, month, item_name, withdrawal_day, description, amount, is_template)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (e.get("id"), e["year"], e["month"], e["item_name"], e.get("withdrawal_day", ""), e.get("description", ""), e["amount"], e.get("is_template", 0)))

        for s in data.get("savings_investments", []):
            cur.execute("""
                INSERT INTO savings_investments (id, year, month, item_name, payment_day, description, amount, is_template)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (s.get("id"), s["year"], s["month"], s["item_name"], s.get("payment_day", ""), s.get("description", ""), s["amount"], s.get("is_template", 0)))

        for v in data.get("variable_expenses", []):
            cur.execute("""
                INSERT INTO variable_expenses (id, transaction_date, category_id, title, payment_method_id, amount, memo, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (v.get("id"), v["transaction_date"], v["category_id"], v["title"], v["payment_method_id"], v["amount"], v.get("memo", ""), v.get("created_at"), v.get("updated_at")))

    return True

def reset_all_data(keep_default_templates: bool = True) -> str:
    """
    Safely resets all recorded transactions and restores default settings.
    A full safety backup snapshot is created automatically before deleting anything.
    """
    import shutil
    from app.core.config import DB_PATH, BACKUP_DIR
    from app.database.schema import (
        DEFAULT_CATEGORIES,
        DEFAULT_PAYMENTS,
        DEFAULT_FIXED_INCOME_TEMPLATES,
        DEFAULT_FIXED_EXPENSE_TEMPLATES,
        DEFAULT_SAVINGS_TEMPLATES
    )

    # 1. Take a safety snapshot backup
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = BACKUP_DIR / f"account_book_before_reset_{timestamp}.db"
    if DB_PATH.exists():
        shutil.copy2(DB_PATH, backup_file)

    # 2. Reset database tables
    with get_db_cursor() as cur:
        # Clear variable expenses
        cur.execute("DELETE FROM variable_expenses")

        # Clear fixed plans
        cur.execute("DELETE FROM fixed_incomes")
        cur.execute("DELETE FROM fixed_expenses")
        cur.execute("DELETE FROM savings_investments")

        # Clear and reseed categories
        cur.execute("DELETE FROM categories")
        cur.execute("DELETE FROM payment_methods")
        try:
            cur.execute("DELETE FROM sqlite_sequence WHERE name IN ('categories', 'payment_methods', 'variable_expenses')")
        except Exception:
            pass

        for name, color, icon, order in DEFAULT_CATEGORIES:
            cur.execute("""
                INSERT INTO categories (name, color_hex, icon, sort_order)
                VALUES (?, ?, ?, ?)
            """, (name, color, icon, order))

        for name, active in DEFAULT_PAYMENTS:
            cur.execute("""
                INSERT INTO payment_methods (name, is_active)
                VALUES (?, ?)
            """, (name, active))

        # Seed default templates (year=0, month=0, is_template=1)
        if keep_default_templates:
            for item, day, desc, amt in DEFAULT_FIXED_INCOME_TEMPLATES:
                cur.execute("""
                    INSERT INTO fixed_incomes (year, month, item_name, expected_day, description, amount, is_template)
                    VALUES (0, 0, ?, ?, ?, ?, 1)
                """, (item, day, desc, amt))

            for item, day, desc, amt in DEFAULT_FIXED_EXPENSE_TEMPLATES:
                cur.execute("""
                    INSERT INTO fixed_expenses (year, month, item_name, withdrawal_day, description, amount, is_template)
                    VALUES (0, 0, ?, ?, ?, ?, 1)
                """, (item, day, desc, amt))

            for item, day, desc, amt in DEFAULT_SAVINGS_TEMPLATES:
                cur.execute("""
                    INSERT INTO savings_investments (year, month, item_name, payment_day, description, amount, is_template)
                    VALUES (0, 0, ?, ?, ?, ?, 1)
                """, (item, day, desc, amt))

    return backup_file.name

